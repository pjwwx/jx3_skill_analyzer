"""Compact, readable exports for the desktop application.

Intermediate CSVs are consumed only after the full analysis has completed. The
original values and diagnostics remain in JSON, independently of Excel limits.
"""
from __future__ import annotations

import csv
import json
import re
from collections import OrderedDict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE


TABLES = OrderedDict([
    ("技能解析.csv", "技能"), ("关联Buff.csv", "关联Buff"),
    ("NPC.csv", "NPC"), ("交互物品.csv", "交互物品"),
    ("地图预置.csv", "地图对象"), ("官方百科.csv", "官方百科"),
])
INPUT_CSVS = set(TABLES) | {"机关.csv", "地图.csv", "官方DBM.csv", "DBM提示.csv",
                          "关系证据.csv", "资料问题.csv", "错误与警告.csv"}
WARNING_RE = re.compile(r"依赖\s+([^:]+):\s*(.*)", re.S)


def _read_json(path: Path) -> dict:
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8-sig"))


def _read_csv(path: Path) -> list[dict]:
    if not path.exists():
        return []
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def _short_message(message: str) -> str:
    if "diagnostic pseudocode" in message or "反编译结果存在恢复异常" in message:
        return "反编译未能可靠恢复；该依赖的逻辑未参与分析。完整诊断见 JSON/解析结果.json。"
    return message


def _merged_issues(tables: dict, workflow: dict) -> list[dict]:
    grouped: OrderedDict[tuple, dict] = OrderedDict()
    inputs = []
    for row in tables.get("错误与警告.csv", []):
        inputs.append((row.get("级别", "提示"), row.get("项目", "技能解析"),
                       row.get("技能ID", ""), row.get("路径", ""), row.get("信息", "")))
    for row in tables.get("资料问题.csv", []):
        path = row.get("来源文件", row.get("路径", row.get("path", "")))
        source_row = row.get("来源行", row.get("row", ""))
        if source_row:
            path += f"（第 {source_row} 行）"
        inputs.append((row.get("级别", row.get("level", "提示")), "副本资料", "",
                       path, row.get("说明", row.get("信息", row.get("message", "")))))
    for row in workflow.get("issues", []):
        # Script reference failures are already represented in skill issues.
        if row.get("kind") == "script_reference" and any(row.get("path", "") in item[4] for item in inputs):
            continue
        if row.get("kind") == "official_dbm_not_read":
            inputs = [item for item in inputs if item[3] != row.get("path")]
        inputs.append((row.get("level", "提示"), "资料读取", "", row.get("path", ""), row.get("message", "")))
    for level, kind, skill_id, path, message in inputs:
        match = WARNING_RE.search(message)
        if match:
            path, message = match.groups()
        if message.startswith("伤害由版本函数动态计算"):
            path = "对应技能脚本（见受影响技能ID）"
        key = (level, kind, path, message)
        if key not in grouped:
            grouped[key] = {"级别": level, "项目": kind, "路径": path,
                            "说明": _short_message(message), "受影响技能ID": [], "出现次数": 0}
        item = grouped[key]
        item["出现次数"] += 1
        if skill_id and skill_id not in item["受影响技能ID"]:
            item["受影响技能ID"].append(skill_id)
    for item in grouped.values():
        item["受影响技能数量"] = len(item["受影响技能ID"])
        item["受影响技能ID"] = "; ".join(item["受影响技能ID"])
    return list(grouped.values())


def _dbm_table(rules: list[dict], prompts: list[dict]) -> list[dict]:
    def key(row):
        return (row.get("来源文件", row.get("source_path", "")),
                row.get("来源行", row.get("source_row", "")))
    by_rule: dict[tuple, list] = {}
    for prompt in prompts:
        by_rule.setdefault(key(prompt), []).append(prompt)
    result = []
    for rule in rules:
        for prompt in by_rule.pop(key(rule), [{}]):
            result.append({**rule, **prompt})
    for remaining in by_rule.values():
        result.extend(remaining)
    return result


def _rewrite_paths(value, replacements: dict[str, str]):
    if isinstance(value, str):
        for old, new in replacements.items():
            value = value.replace(old, new).replace(old.replace("\\", "/"), new.replace("\\", "/"))
        return value
    if isinstance(value, list):
        return [_rewrite_paths(item, replacements) for item in value]
    if isinstance(value, dict):
        return {_rewrite_paths(key, replacements): _rewrite_paths(item, replacements) for key, item in value.items()}
    return value


def _cell_text(value):
    if isinstance(value, (list, dict)):
        value = json.dumps(value, ensure_ascii=False)
    if value is None:
        return ""
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
        if len(value) > 32767:
            value = value[:32680] + "\n[内容较长，完整值见 JSON/表格数据.json]"
    return value


PRIMARY_FIELDS = {
    "技能": ["技能ID", "技能名称", "伤害类型", "是否穿刺", "是否穿透", "技能类型", "技能形状",
             "技能释放方式", "技能读条时间(秒)", "技能引导时间(秒)", "技能作用半径", "最大释放距离",
             "目标数量上限", "技能宽度", "技能高度", "技能角度", "最小释放距离", "技能保护半径",
             "技能控制效果", "技能是否可打断", "各等级基础伤害", "各等级伤害浮动", "UI描述"],
    "关联Buff": ["来源技能ID", "来源技能名称", "BuffID", "Buff名称",
                 "关联方式", "配置等级", "最大层数", "间隔时间(秒)", "估算持续时间(秒)", "UI描述", "属性概览"],
    "NPC": ["NPC模板ID", "NPC名称", "称号", "等级原值", "最大生命值原值", "生命计数字段原值", "实战血量",
            "技能槽原值", "AI关联状态", "NPC脚本状态", "尸体交互物品名称", "尸体关联状态", "关联副本名称", "读取依据"],
    "交互物品": ["交互物品模板ID", "交互物品名称", "关联副本名称", "名称来源", "可选中原值", "可拾取原值",
                 "逐人操作字段原值", "准备帧原值", "读条文字", "准备时间说明", "脚本状态", "具体交互", "关联依据", "交互限制"],
    "地图对象": ["预置类别", "模板ID", "模板名称", "预置昵称", "副本名称", "X原值", "Y原值", "Z原值", "模板关联状态", "证据说明"],
    "官方百科": ["百科BOSS名称", "阶段", "技能文案名称", "技能文案说明", "展示NPC名称", "副本名称", "百科地图索引状态", "文案关联状态", "百科关系说明"],
    "官方DBM": ["副本名称", "规则ID", "监测对象类型", "目标ID原值", "提示文字", "触发条件", "触发后秒数原值", "时间说明"],
    "问题": ["级别", "项目", "说明", "受影响技能数量", "出现次数", "路径", "受影响技能ID"],
}
DETAIL_FIELDS = {
    "技能": {"技能名称(配置)", "技能名称(UI)", "UI短描述", "UI简述", "技能读条帧", "技能引导帧",
             "属性类型(原始)", "解析到的依赖函数", "基础伤害表达式", "伤害浮动表达式"},
    "关联Buff": {"Buff名称(配置)", "Buff名称(UI)", "BuffID表达式", "配置等级表达式", "引用等级表达式", "计数(Count)", "间隔帧(Interval)", "Buff脚本"},
    "NPC": {"AI类型ID", "AI脚本", "NPC脚本", "尸体交互物品模板ID", "关联副本ID", "模板所属地图"},
    "交互物品": {"关联副本ID", "模板所属地图", "交互脚本", "等级原值", "分类ID原值"},
    "地图对象": {"副本ID", "来源记录序号"},
    "官方百科": {"副本ID", "百科BOSS索引", "展示NPC模板ID", "技能文案ID"},
    "问题": {"受影响技能ID"},
}


def _column_width(title: str, field: str) -> float:
    if title == "概览":
        return {"项目": 20, "内容": 54, "说明": 68}[field]
    explicit = {"技能名称": 46, "来源技能名称": 32, "NPC名称": 26, "称号": 24, "交互物品名称": 34,
                "技能名称(配置)": 44, "技能名称(UI)": 24, "Buff名称(UI)": 24, "Buff名称(配置)": 24,
                "Buff名称": 40, "预置昵称": 26, "模板名称": 32,
                "伤害类型": 25, "技能类型": 14, "技能形状": 13, "技能释放方式": 15,
                "是否穿刺": 11, "是否穿透": 11, "技能读条时间(秒)": 15, "技能引导时间(秒)": 15,
                "实战血量": 27, "最大生命值原值": 20, "生命计数字段原值": 20, "技能槽原值": 40,
                "说明": 68 if title == "问题" else 48, "路径": 55, "受影响技能ID": 70,
                "受影响技能数量": 16, "出现次数": 12, "技能文案说明": 64, "技能文案名称": 30,
                "预告提示": 52, "提示文字": 48, "各等级基础伤害": 38, "各等级伤害浮动": 34,
                "UI描述": 52, "解析说明": 48, "快照ID": 72, "X原值": 15, "Y原值": 15, "Z原值": 13}
    if field in explicit:
        return explicit[field]
    if field.endswith("ID") or field.endswith("索引") or "来源行" in field or field == "来源记录序号":
        return max(10, min(18, len(field) * 2 + 2))
    if "状态" in field:
        return 22
    if any(word in field for word in ("脚本", "文件", "表达式")):
        return 44
    if any(word in field for word in ("描述", "说明", "概览", "文案")):
        return 48
    if any(word in field for word in ("名称", "副本", "引用来源")):
        return 26
    return max(12, min(20, len(field) * 2 + 2))


def _add_table(book: Workbook, title: str, rows: list[dict]):
    if not rows:
        return None
    sheet = book.create_sheet(title)
    rows = [dict(row) for row in rows]
    if title == "技能":
        for row in rows:
            row["技能名称"] = row.get("技能名称(UI)") or row.get("技能名称(配置)") or ""
    elif title == "关联Buff":
        for row in rows:
            row["Buff名称"] = row.get("Buff名称(UI)") or row.get("Buff名称(配置)") or ""
    original_fields = list(dict.fromkeys(key for row in rows for key in row))
    details = set(DETAIL_FIELDS.get(title, set()))
    details.update(field for field in original_fields if field in {"快照ID", "来源文件", "来源行", "文案来源文件", "文案来源行"})
    ordered = [field for field in PRIMARY_FIELDS.get(title, []) if field in original_fields]
    ordered.extend(field for field in original_fields if field not in ordered and field not in details)
    fields = ordered + [field for field in original_fields if field not in ordered]
    sheet.append(fields)
    for row in rows:
        sheet.append([_cell_text(row.get(key, "")) for key in fields])
    # IDs and source expressions stay text. Formula-looking source strings must
    # never be interpreted as Excel formulas.
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.font = Font(name="Microsoft YaHei", size=11, color="243B53")
            cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=False)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="294860")
        cell.font = Font(name="Microsoft YaHei", size=11, color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
        cell.border = Border(right=Side(style="thin", color="FFFFFF"))
    sheet.freeze_panes = "D2" if title == "地图对象" else "C2" if len(fields) > 3 else "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.zoomScale = 95
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.outlinePr.summaryRight = True
    sheet.sheet_format.defaultRowHeight = 24
    sheet.row_dimensions[1].height = 38
    for index, field in enumerate(fields, 1):
        column = sheet.column_dimensions[get_column_letter(index)]
        column.width = _column_width(title, field)
        if field in details:
            column.hidden = True
            column.outlineLevel = 1
    if details:
        sheet.column_dimensions[get_column_letter(len(fields) + 1)].collapsed = True
    for index in range(2, sheet.max_row + 1):
        if title == "概览":
            # Only the three visible summary cells affect this row. Counts stay
            # compact, while longer notes get room for a few readable lines.
            lines = max(sum(max(1, (sum(2 if ord(char) > 127 else 1 for char in line) +
                                   _column_width(title, field) - 1) // _column_width(title, field))
                            for line in str(cell.value or "").split("\n"))
                        for cell, field in zip(sheet[index], fields))
            sheet.row_dimensions[index].height = min(60, max(28, lines * 16 + 8))
        else:
            sheet.row_dimensions[index].height = 44 if title == "问题" else 24
        for cell, field in zip(sheet[index], fields):
            cell.border = Border(bottom=Side(style="hair", color="E9EEF3"))
            if index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F6F8FB")
            if isinstance(cell.value, (int, float)):
                cell.alignment = Alignment(vertical="center", horizontal="right", wrap_text=False)
            if title in {"概览", "问题"}:
                cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=field not in {"受影响技能ID", "路径"})
            elif isinstance(cell.value, str) and "\n" in cell.value:
                cell.alignment = Alignment(vertical="top", horizontal="left", wrap_text=False)
            if field in {"技能名称", "Buff名称", "NPC名称", "交互物品名称", "百科BOSS名称"}:
                cell.font = Font(name="Microsoft YaHei", size=11, color="243B53", bold=True)
            elif field in {"是否穿刺", "是否穿透"} and cell.value:
                cell.font = Font(name="Microsoft YaHei", size=11, color="16705B" if field == "是否穿刺" else "245AA0", bold=True)
    return sheet


def _overview(payload: dict, resources: dict, tables: dict, issues: list[dict]) -> list[dict]:
    rows = []
    def add(item, value, note=""):
        rows.append({"项目": item, "内容": value, "说明": note})
    names = payload.get("selected_dungeons", [])
    add("所选副本", "; ".join(str(item) for item in names))
    add("地图ID", "; ".join(payload.get("selected_map_ids", resources.get("selected_map_ids", []))))
    add("版本", payload.get("version", ""))
    is_v4 = payload.get("package_source") == "local PakV4"
    add("数据来源", "本机 V4 包" if is_v4 else "手动解包资料", "按所选副本引用读取。" if is_v4 else "使用所选本地资料和技能ID。")
    add("技能", len(tables.get("技能解析.csv", [])))
    add("未完成技能", max(0, len(payload.get("requested_skill_ids", [])) - len(tables.get("技能解析.csv", []))))
    add("关联Buff", len(tables.get("关联Buff.csv", [])))
    add("NPC记录", len(tables.get("NPC.csv", [])), "同一模板可有多条配置记录；血量列为配置原值。")
    add("交互物品", len(tables.get("交互物品.csv", tables.get("机关.csv", []))))
    add("地图对象", len(tables.get("地图预置.csv", [])))
    add("官方百科Boss", len(resources.get("encyclopaedia", [])),
        f"展开为 {len(tables.get('官方百科.csv', []))} 条文案记录。")
    dbm_rows = len(tables.get("官方DBM.csv", []))
    add("官方DBM有效规则", resources.get("counts", {}).get("active_dbm_rules", dbm_rows),
        f"共 {dbm_rows} 条原始规则记录。没有读取到资料时不创建空工作表；具体缺失见‘问题’。")
    add("问题", len(issues), "重复问题已合并，列出受影响技能。")
    add("阅读方式", "先看概览，再看各工作表", "表头可筛选；JSON保存完整字段和诊断，原始资料保存提取和反编译文件。")
    add("长文字", "默认紧凑行高", "点击单元格可在公式栏查看完整内容，也可手动增加行高。")
    add("辅助字段", "折叠于表格右侧", "点击列上方的加号展开原始名称、帧数、脚本和来源字段。")
    add("伤害标识", "穿刺与穿透分别列出", "SkillPuncture 为穿刺伤害，SkillPenetration 为穿透。")
    if tables.get("地图.csv"):
        for index, item in enumerate(tables["地图.csv"], 1):
            # Only the primary identity fields are needed in the overview.
            values = {key: item.get(key, "") for key in ("副本ID", "地图ID", "副本名称", "显示名称", "人数上限原值", "人数上限", "ID", "Name", "DisplayName", "MaxPlayerCount") if item.get(key, "")}
            add(f"地图 {index}", "; ".join(f"{key}: {value}" for key, value in values.items()))
    for note in resources.get("limitations", []):
        add("资料范围", str(note).replace("机关", "交互物品"))
    return rows


def finalize_output_layout(output_dir: Path | str, result=None) -> dict:
    """Finalize this run after all producers have finished; safe to call twice."""
    root = Path(output_dir).resolve()
    workbook_path = root / "副本解析.xlsx"
    json_dir, raw_dir = root / "JSON", root / "原始资料"
    if workbook_path.exists() and not (root / "技能解析.csv").exists():
        info = {"workbook": workbook_path, "json_dir": json_dir, "raw_dir": raw_dir,
                "json_file": json_dir / "解析结果.json", "old_to_new_paths": {}}
        if result is not None:
            _update_result(result, info)
        return info
    tables = {name: _read_csv(root / name) for name in sorted(INPUT_CSVS) if (root / name).exists()}
    if "机关.csv" in tables and "交互物品.csv" not in tables:
        tables["交互物品.csv"] = tables["机关.csv"]
    payload = _read_json(root / "解析结果.json")
    run_inputs = {name: (root / name).read_text(encoding="utf-8-sig") for name in ("所选副本.txt", "技能ID.txt") if (root / name).exists()}
    resources = _read_json(root / "副本资料.json")
    issues = _merged_issues(tables, _read_json(root / "资源读取状态.json"))
    planned_moves = {str(root / name): str(raw_dir / name) for name in ("解包原文件", "反编译脚本")}
    issues = _rewrite_paths(issues, planned_moves)
    for item in issues:
        for prefix in (str(raw_dir) + "\\", str(raw_dir) + "/"):
            if item["路径"].startswith(prefix):
                item["路径"] = "原始资料/" + item["路径"][len(prefix):].replace("\\", "/")
    readable_tables = _rewrite_paths(tables, planned_moves)
    book = Workbook()
    book.remove(book.active)
    _add_table(book, "概览", _overview(payload, resources, tables, issues))
    for filename, sheet_name in TABLES.items():
        rows = readable_tables.get(filename, [])
        if sheet_name == "技能":
            rows = [dict(row) for row in rows]
            for row in rows:
                message = row.get("解析说明", "")
                if "diagnostic pseudocode" in message:
                    row["解析说明"] = "依赖反编译不完整，详见‘问题’工作表。"
                elif "共享脚本整体仍不可靠" in message:
                    row["解析说明"] = re.sub(r"依赖\s+[^;]*Skill\.lh:[^;]*", "Skill.lh 部分恢复，详见‘问题’。", message)
        _add_table(book, sheet_name, rows)
    _add_table(book, "官方DBM", _dbm_table(tables.get("官方DBM.csv", []), tables.get("DBM提示.csv", [])))
    _add_table(book, "问题", issues or [{"级别": "", "项目": "", "路径": "", "说明": "本次未记录问题。"}])
    generated_sheets = book.sheetnames
    temporary = root / ".副本解析.tmp.xlsx"
    book.save(temporary)
    book.close()
    # Keep everything until the workbook has been saved successfully.
    json_dir.mkdir(exist_ok=True)
    raw_dir.mkdir(exist_ok=True)
    replacements = {}
    for name in ("解包原文件", "反编译脚本"):
        source, target = root / name, raw_dir / name
        if source.exists():
            if target.exists():
                raise FileExistsError(f"整理结果目标已存在：{target}")
            source.rename(target)
            replacements[str(source)] = str(target)
    json_files = list(root.glob("*.json"))
    replacements.update({str(path): str(json_dir / path.name) for path in json_files})
    for name in INPUT_CSVS:
        replacements[str(root / name)] = str(json_dir / "表格数据.json") if name == "关系证据.csv" else str(workbook_path)
    for path in json_files:
        target = json_dir / path.name
        data = _rewrite_paths(_read_json(path), replacements)
        if path.name == "解析结果.json":
            if run_inputs:
                data["run_inputs"] = run_inputs
            sheet_map = {"skills": "技能", "buffs": "关联Buff", "issues": "问题",
                         "npc_templates": "NPC", "doodad_templates": "交互物品",
                         "placements": "地图对象", "encyclopaedia": "官方百科",
                         "dbm_rules": "官方DBM", "dbm_prompts": "官方DBM", "maps": "概览"}
            data["output_layout"] = {"version": 1, "workbook": str(workbook_path),
                                     "sheets": {key: title for key, title in sheet_map.items() if title in generated_sheets},
                                     "complete_tables": str(json_dir / "表格数据.json")}
        target.write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding="utf-8-sig")
        path.unlink()
    (json_dir / "表格数据.json").write_text(json.dumps(_rewrite_paths(tables, replacements), ensure_ascii=False, indent=2), encoding="utf-8-sig")
    (json_dir / "问题汇总.json").write_text(json.dumps(issues, ensure_ascii=False, indent=2), encoding="utf-8-sig")
    for name in ("所选副本.txt", "技能ID.txt"):
        source = root / name
        if source.exists():
            source.unlink()
    for name in INPUT_CSVS:
        (root / name).unlink(missing_ok=True)
    temporary.replace(workbook_path)
    info = {"workbook": workbook_path, "json_dir": json_dir, "raw_dir": raw_dir,
            "json_file": json_dir / "解析结果.json", "old_to_new_paths": replacements}
    if result is not None:
        _update_result(result, info)
    return info


def _update_result(result, info: dict):
    result.skill_csv = result.buff_csv = result.issue_csv = info["workbook"]
    result.json_file = info["json_file"]
    result.workbook = info["workbook"]
    result.output_layout_finalized = True
